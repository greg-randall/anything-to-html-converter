"""Extracted images: each output's own media folder, links that resolve from the HTML, and missing images.

Pandoc runs only on the tiny documents written into the test's temp folder.
"""
import re
import sys
import urllib.parse

import docx
from PIL import Image
from test_end_to_end import ALT_TEXT, prompt_content, strip_ansi


def make_docx(path, image_path):
    """Write a .docx with one paragraph and one picture (python-docx stores it as word/media/image1.png)."""
    document = docx.Document()
    document.add_paragraph(f"Document {path.stem}")
    document.add_picture(str(image_path))
    document.save(str(path))


def image_links(markdown_text):
    return re.findall(r'!\[[^\]]*\]\(([^)]*)\)', markdown_text)


def html_image_files(html_path):
    """The file each <img src> in the HTML points to, resolved from the HTML file's folder."""
    html = html_path.read_text(encoding='utf-8')
    sources = re.findall(r'<img src="([^"]*)"', html)
    return [html_path.parent / urllib.parse.unquote(src) for src in sources]


def test_alt_text_is_not_shown_as_a_visible_caption(api):
    html = api.convert_markdown_to_html('Text before.\n\n![Two people holding signs](pic.png)\n\nText after.\n')
    assert 'alt="Two people holding signs"' in html
    assert '<figure' not in html
    assert '<figcaption' not in html
    body = re.sub(r'<[^>]+>', '', html.split('<body>', 1)[1])
    assert 'Two people holding signs' not in body


def test_media_folder_name(api):
    assert api.media_folder_name('Example Document_improved.html') == (
        'Example Document_improved_media')
    assert api.media_folder_name('/some/dir/out.htm') == 'out_media'


def test_two_docx_inputs_in_one_folder_keep_separate_images(api, tmp_path):
    for name, color in (('first doc', 'red'), ('second doc', 'blue')):
        Image.new('RGB', (4, 4), color).save(tmp_path / f'{color}.png', format='PNG')
        make_docx(tmp_path / f'{name}.docx', tmp_path / f'{color}.png')

    results = {}
    for name in ('first doc', 'second doc'):
        html_output = tmp_path / f'{name}_improved.html'
        markdown = api.convert_to_markdown(str(tmp_path / f'{name}.docx'), str(tmp_path / f'{name}.md'),
                                           str(html_output))
        # The .docx's internal media/ level is dropped: images sit directly in the output's own folder
        assert image_links(markdown) == [f'{name}_improved_media/image1.png']
        html_output.write_text(api.convert_markdown_to_html(markdown), encoding='utf-8')
        results[name] = html_output

    # Each HTML still shows its own picture after the second conversion
    for name, color in (('first doc', 'red'), ('second doc', 'blue')):
        files = html_image_files(results[name])
        assert files == [tmp_path / f'{name}_improved_media' / 'image1.png']
        with Image.open(files[0]) as image:
            assert image.convert('RGB').getpixel((0, 0)) == Image.new('RGB', (1, 1), color).getpixel((0, 0))
    assert not (tmp_path / 'media').exists()


def test_markdown_input_with_output_in_another_folder_finds_its_images(api, monkeypatch, capsys, tmp_path):
    source = tmp_path / 'source'
    (source / 'pics').mkdir(parents=True)
    Image.new('RGB', (4, 4), 'green').save(source / 'pics' / 'logo.png', format='PNG')
    (source / 'news.md').write_text("# News\n\nWords here.\n\n![](pics/logo.png)\n\nMore words.\n", encoding='utf-8')
    out_dir = tmp_path / 'out'
    out_dir.mkdir()
    # Run from a third folder, so neither the input's nor the output's folder is the working directory
    elsewhere = tmp_path / 'elsewhere'
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)
    monkeypatch.delenv('LLM_BASE_URL', raising=False)
    monkeypatch.setenv('LLM_API', 'anthropic')
    monkeypatch.setenv('LLM_API_KEY', 'fake-key')
    monkeypatch.setenv('FORMAT_MODEL', 'test-format-deployment')
    monkeypatch.setenv('FORMAT_MAX_TOKENS', '1000')
    monkeypatch.setenv('ALT_TEXT_MODEL', 'test-alt-deployment')
    monkeypatch.setenv('ALT_TEXT_MAX_TOKENS', '200')
    monkeypatch.setenv('ALT_TEXT_CONTEXT_WORDS', '5')
    monkeypatch.setenv('ALT_TEXT_CONCURRENCY', '2')
    api.use_prices(monkeypatch, "{}\n", tmp_path / 'prices.yaml')
    calls = api.fake_llms(monkeypatch, format_reply=lambda system, user: prompt_content(user),
                          alt_text_reply=lambda system, content: ALT_TEXT)

    monkeypatch.setattr(sys, 'argv', ['anything-to-anything.py', str(source / 'news.md'),
                                      '--output', str(out_dir / 'result.html')])
    api.main()
    out = strip_ansi(capsys.readouterr().out)

    assert "could not find" not in out
    improved_md = (out_dir / 'result.md').read_text(encoding='utf-8')
    assert image_links(improved_md) == ['result_media/logo.png']
    assert html_image_files(out_dir / 'result.html') == [out_dir / 'result_media' / 'logo.png']
    assert (out_dir / 'result_media' / 'logo.png').read_bytes() == (
        source / 'pics' / 'logo.png').read_bytes()
    # Alt text generation found the extracted file (an image content block was sent)
    # The calls run in parallel, so they can be recorded in either order
    assert sorted(c['kind'] for c in calls) == ['alt_text', 'format']
    alt_call = next(c for c in calls if c['kind'] == 'alt_text')
    assert any(block.get('type') == 'image' for block in alt_call['user'])
    assert f"![{ALT_TEXT}](result_media/logo.png)" in improved_md


def test_missing_image_gets_a_loud_warning_naming_it(api, capsys, tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    Image.new('RGB', (4, 4), 'green').save(source / 'found.png', format='PNG')
    (source / 'page.html').write_text('<p>Hi</p><img src="found.png"><img src="gone/missing one.png" alt="Lost">'
                                      '<img src="also-missing.jpg">\n', encoding='utf-8')
    out_dir = tmp_path / 'out'
    out_dir.mkdir()
    markdown = api.convert_to_markdown(str(source / 'page.html'), str(out_dir / 'page.md'),
                                       str(out_dir / 'page_improved.html'))
    out = strip_ansi(capsys.readouterr().out)

    assert ("\nWARNING: pandoc could not find 2 image(s) in the input. Each was replaced with a text placeholder, "
            "so it will be missing from the output:\n"
            # pandoc records the path URL-encoded
            f"  - gone/missing%20one.png (relative paths are looked up in {source})\n"
            f"  - also-missing.jpg (relative paths are looked up in {source})\n") in out
    assert image_links(markdown) == ['page_improved_media/found.png']
    assert (out_dir / 'page_improved_media' / 'found.png').is_file()


def write_files(base, paths):
    for path in paths:
        target = base.joinpath(*path.split('/'))
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(path.encode())


def test_flatten_drops_the_single_wrapper_folder_and_rewrites_image_links(api, tmp_path):
    write_files(tmp_path / 'out_media.extracting', ['media/image1.png', 'media/sub/image2.png'])
    markdown = ("![](out_media.extracting/media/image1.png){width=\"1in\"}\n\n"
                "![a](out_media.extracting/media/sub/image2.png)\n\n"
                "[a link](out_media.extracting/media/image1.png)\n")
    result = api.flatten_extracted_media(markdown, str(tmp_path), 'out_media.extracting', 'out_media')
    assert result == ("![](out_media/image1.png){width=\"1in\"}\n\n"
                      "![a](out_media/sub/image2.png)\n\n"
                      "[a link](out_media.extracting/media/image1.png)\n")
    assert (tmp_path / 'out_media' / 'image1.png').read_bytes() == b'media/image1.png'
    assert (tmp_path / 'out_media' / 'sub' / 'image2.png').read_bytes() == b'media/sub/image2.png'
    assert not (tmp_path / 'out_media.extracting').exists()


def test_flatten_keeps_folders_when_images_are_not_under_one_wrapper(api, tmp_path):
    write_files(tmp_path / 'out_media.extracting', ['pics/a.png', 'logos/b.png'])
    markdown = "![](out_media.extracting/pics/a.png)\n\n![](out_media.extracting/logos/b.png)\n"
    result = api.flatten_extracted_media(markdown, str(tmp_path), 'out_media.extracting', 'out_media')
    assert result == "![](out_media/pics/a.png)\n\n![](out_media/logos/b.png)\n"
    assert (tmp_path / 'out_media' / 'pics' / 'a.png').exists()
    assert (tmp_path / 'out_media' / 'logos' / 'b.png').exists()


def test_flatten_with_nothing_extracted_changes_nothing(api, tmp_path):
    assert api.flatten_extracted_media("No images.\n", str(tmp_path), 'x.extracting', 'x') == "No images.\n"
    assert list(tmp_path.iterdir()) == []


def test_rerunning_the_same_docx_replaces_its_images_and_leaves_no_staging_folder(api, tmp_path):
    Image.new('RGB', (4, 4), 'red').save(tmp_path / 'red.png', format='PNG')
    Image.new('RGB', (4, 4), 'blue').save(tmp_path / 'blue.png', format='PNG')
    html_output = tmp_path / 'doc_improved.html'
    for color in ('red', 'blue'):
        make_docx(tmp_path / 'doc.docx', tmp_path / f'{color}.png')
        markdown = api.convert_to_markdown(str(tmp_path / 'doc.docx'), str(tmp_path / 'doc.md'), str(html_output))
    assert image_links(markdown) == ['doc_improved_media/image1.png']
    assert sorted(p.name for p in (tmp_path / 'doc_improved_media').iterdir()) == ['image1.png']
    assert (tmp_path / 'doc_improved_media' / 'image1.png').read_bytes() == (tmp_path / 'blue.png').read_bytes()
    assert not (tmp_path / 'doc_improved_media.extracting').exists()
